"""Input adapter: Validation_Rules_Master.xlsx -> engine template CSVs.

This is the ONE place that knows the layout of the business-authored master
workbook. It translates that workbook into the flat CSV contract the Option B
engine consumes. Business analysts edit the .xlsx; engineers own this adapter
and the engine.

INPUT
    rules/Validation_Rules_Master.xlsx
      * Four "table tabs" (Contract, Address, Roles, Coverage) — one row per
        (field, rule) pair. Columns:
          Field Name | Data Type | Length | Precision | Date Format |
          Is Key Field | Rule Type | Severity | Allowable Codes |
          Allowable Description | Description / Notes
      * One "Cross_References" tab — one row per cross-reference rule. Columns:
          XRef ID | Table Name | Target Field | Rule Type | Severity |
          Precondition Field | Precondition Operator | Precondition Value |
          Expected Result | Compare Table | Compare Field | Join Child Fields |
          Join Parent Fields | Error Message | Active | Notes

OUTPUT (written to templates/v1/)
    validation_rules.csv    — field-level rules (schema: VALIDATION_RULES_*)
    cross_references.csv     — cross-reference rules (schema: CROSS_REFERENCES_*)
    table_dependencies.csv   — parent/child join map (regenerated, stable)
    _reconciliation_report.txt — counts + anything carried-but-not-evaluated

DESIGN NOTES / EXTENSIBILITY
    * Rule types are authored by data engineers in the sheet and copied through
      verbatim. This adapter does NOT infer rule types from prose.
    * Every row is preserved. Rows whose rule_type is not yet understood by the
      engine are still written to the CSV and listed in the reconciliation
      report, so nothing is silently dropped and coverage can grow over time by
      adding a handler (no adapter change needed for a newly-supported type).
    * To onboard a NEW table tab: add its name->code entry to TABLE_TAB_TO_CODE.
    * To change the output location or version: edit OUTPUT_DIR.
"""
import csv
import os
import sys

import openpyxl

WORKSPACE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MASTER_PATH = os.path.join(WORKSPACE, "rules", "Validation_Rules_Master.xlsx")
OUTPUT_DIR = os.path.join(WORKSPACE, "templates", "v1")

# Make the engine's registered types importable for reconciliation.
sys.path.insert(0, WORKSPACE)
from option_b.config.schema_definitions import (  # noqa: E402
    REGISTERED_RULE_TYPES,
    REGISTERED_CROSS_REFERENCE_TYPES,
    VALIDATION_RULES_REQUIRED_COLUMNS,
    VALIDATION_RULES_OPTIONAL_COLUMNS,
    CROSS_REFERENCES_REQUIRED_COLUMNS,
    CROSS_REFERENCES_OPTIONAL_COLUMNS,
)
from option_b.engine.rule_compiler import (  # noqa: E402
    CROSS_REFERENCE_HANDLER_REGISTRY,
    CROSS_REFERENCE_JOIN_TYPES,
    CROSS_REFERENCE_AGGREGATE_TYPES,
)

# Business tab name -> engine table code (field prefixes use these codes).
TABLE_TAB_TO_CODE = {
    "Contract": "CNT",
    "Address": "ADD",
    "Roles": "ROL",
    "Coverage": "CVG",
}

# Free-text table references in the sheet -> engine table code. The sheet uses
# mixed casing/synonyms (CONTRACT, Contract, ADDRESS, ROLE, Roles, ...).
TABLE_NAME_ALIASES = {
    "contract": "CNT",
    "cnt": "CNT",
    "address": "ADD",
    "add": "ADD",
    "roles": "ROL",
    "role": "ROL",
    "rol": "ROL",
    "coverage": "CVG",
    "cvg": "CVG",
}

# Output column order (kept explicit so CSV diffs are stable and reviewable).
VALIDATION_RULES_COLUMNS = (
    VALIDATION_RULES_REQUIRED_COLUMNS + VALIDATION_RULES_OPTIONAL_COLUMNS
)
CROSS_REFERENCES_COLUMNS = (
    CROSS_REFERENCES_REQUIRED_COLUMNS + CROSS_REFERENCES_OPTIONAL_COLUMNS
)

# Stable parent/child dependency map for the current 4-table scope.
TABLE_DEPENDENCIES = [
    {"child_table": "CVG", "parent_table": "CNT",
     "join_fields_child": "CVG_COMPANY_CODE|CVG_POLICY_NUMBER",
     "join_fields_parent": "CNT_COMPANY_CODE|CNT_POLICY_NUMBER"},
    {"child_table": "ROL", "parent_table": "CNT",
     "join_fields_child": "ROL_COMPANY_CODE|ROL_POLICY_NUMBER",
     "join_fields_parent": "CNT_COMPANY_CODE|CNT_POLICY_NUMBER"},
    {"child_table": "ADD", "parent_table": "CNT",
     "join_fields_child": "ADD_COMPANY_CODE|ADD_POLICY_NUMBER",
     "join_fields_parent": "CNT_COMPANY_CODE|CNT_POLICY_NUMBER"},
]


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _clean(value) -> str:
    """Normalize a cell value to a trimmed string ('' for None/NBSP/'N/A')."""
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ").strip()
    if text.upper() in ("N/A", "NA", "NONE"):
        return ""
    return text


def _normalize_pipe_fields(value: str) -> str:
    """Trim each token in a pipe/tab-delimited field list, drop blanks."""
    if not value:
        return ""
    tokens = [t.strip() for t in value.replace("\t", "|").split("|")]
    return "|".join(t for t in tokens if t)


def _table_code(name: str) -> str:
    """Map a free-text table reference to the engine table code (best effort)."""
    key = _clean(name).lower()
    return TABLE_NAME_ALIASES.get(key, _clean(name).upper())


def _header_index(ws) -> dict:
    """Return {header_text: 1-based column index} from row 1."""
    idx = {}
    for c in range(1, ws.max_column + 1):
        h = ws.cell(row=1, column=c).value
        if h is not None and str(h).strip():
            idx[str(h).strip()] = c
    return idx


# ---------------------------------------------------------------------------
# Field-level rules (the four table tabs)
# ---------------------------------------------------------------------------

def build_validation_rules(wb):
    """Read the four table tabs into validation_rules rows.

    Returns (rows, stats, unknown_types) where:
      rows          : list[dict] matching VALIDATION_RULES_COLUMNS
      stats         : {table_code: {"rows": n, "fields": n}}
      unknown_types : list[(table_code, field, rule_type)] not in registry
    """
    rows = []
    stats = {}
    unknown = []

    for tab_name, code in TABLE_TAB_TO_CODE.items():
        if tab_name not in wb.sheetnames:
            print(f"  WARNING: tab '{tab_name}' not found — skipping")
            continue
        ws = wb[tab_name]
        idx = _header_index(ws)

        def col(name):
            return idx.get(name)

        seq = 0
        fields_seen = set()
        for r in range(2, ws.max_row + 1):
            field = _clean(ws.cell(row=r, column=col("Field Name")).value) if col("Field Name") else ""
            rule_type = _clean(ws.cell(row=r, column=col("Rule Type")).value) if col("Rule Type") else ""
            if not field or not rule_type:
                continue

            seq += 1
            fields_seen.add(field)

            severity = _clean(ws.cell(row=r, column=col("Severity")).value) if col("Severity") else "SEVERE"
            length = _clean(ws.cell(row=r, column=col("Length")).value) if col("Length") else ""
            precision = _clean(ws.cell(row=r, column=col("Precision")).value) if col("Precision") else ""
            allow_codes = _clean(ws.cell(row=r, column=col("Allowable Codes")).value) if col("Allowable Codes") else ""
            notes = _clean(ws.cell(row=r, column=col("Description / Notes")).value) if col("Description / Notes") else ""

            param_1, param_2 = _params_for_rule(rule_type, length, precision)
            allowed_values = _normalize_pipe_fields(allow_codes.replace(",", "|"))

            if rule_type not in REGISTERED_RULE_TYPES:
                unknown.append((code, field, rule_type))

            rows.append({
                "rule_id": f"{code}_{seq:03d}",
                "target_table": code,
                "target_field": field,
                "rule_type": rule_type,
                "severity": severity or "SEVERE",
                "sequence": str(seq),
                "param_1": param_1,
                "param_2": param_2,
                "allowed_values": allowed_values,
                "compare_field": "",
                "compare_table": "",
                "compare_column": "",
                "precondition_field": "",
                "precondition_operator": "",
                "precondition_value": "",
                "description": notes.replace("\n", " ").replace("\r", " ")[:300],
            })

        stats[code] = {"rows": seq, "fields": len(fields_seen)}

    return rows, stats, unknown


def _params_for_rule(rule_type: str, length: str, precision: str):
    """Derive (param_1, param_2) for length/precision-driven rule types.

    Length/precision come straight from the sheet's dedicated columns, so no
    prose parsing is needed. Rule types that don't use params get ("","").
    """
    length = length if length.isdigit() else (length if length else "")
    if rule_type in ("exact_length", "max_length", "exact_length_or_blank", "project_defined_value"):
        return (length, "")
    if rule_type == "numeric_decimal":
        return (length, precision)
    # numeric_range / min_value bounds live in the cross-reference/expected
    # metadata or are seeded by DEs later; leave params empty here.
    return ("", "")


# ---------------------------------------------------------------------------
# Cross-reference rules (the Cross_References tab)
# ---------------------------------------------------------------------------

def build_cross_references(wb):
    """Read the Cross_References tab into cross_references rows.

    Returns (rows, stats, not_evaluated) where:
      rows          : list[dict] matching CROSS_REFERENCES_COLUMNS
      stats         : {rule_type: count}
      not_evaluated : list[(xref_id, rule_type)] the engine cannot yet evaluate
    """
    if "Cross_References" not in wb.sheetnames:
        print("  WARNING: 'Cross_References' tab not found — no cross-refs emitted")
        return [], {}, []

    ws = wb["Cross_References"]
    idx = _header_index(ws)

    def col(name):
        return idx.get(name)

    # Cross-reference types the engine has a handler for (column, join, or aggregate).
    evaluable = (
        set(CROSS_REFERENCE_HANDLER_REGISTRY.keys())
        | set(CROSS_REFERENCE_JOIN_TYPES)
        | set(CROSS_REFERENCE_AGGREGATE_TYPES)
    )

    rows = []
    stats = {}
    not_evaluated = []

    for r in range(2, ws.max_row + 1):
        xref_id = _clean(ws.cell(row=r, column=col("XRef ID")).value) if col("XRef ID") else ""
        rule_type = _clean(ws.cell(row=r, column=col("Rule Type")).value) if col("Rule Type") else ""
        if not xref_id or not rule_type:
            continue

        stats[rule_type] = stats.get(rule_type, 0) + 1

        table_name = _table_code(ws.cell(row=r, column=col("Table Name")).value) if col("Table Name") else ""
        target_field = _clean(ws.cell(row=r, column=col("Target Field")).value) if col("Target Field") else ""
        severity = _clean(ws.cell(row=r, column=col("Severity")).value) if col("Severity") else "SEVERE"
        pre_field = _clean(ws.cell(row=r, column=col("Precondition Field")).value) if col("Precondition Field") else ""
        pre_op = _clean(ws.cell(row=r, column=col("Precondition Operator")).value) if col("Precondition Operator") else ""
        pre_val = _clean(ws.cell(row=r, column=col("Precondition Value")).value) if col("Precondition Value") else ""
        pre_field2 = _clean(ws.cell(row=r, column=col("Precondition Field 2")).value) if col("Precondition Field 2") else ""
        pre_op2 = _clean(ws.cell(row=r, column=col("Precondition Operator 2")).value) if col("Precondition Operator 2") else ""
        pre_val2 = _clean(ws.cell(row=r, column=col("Precondition Value 2")).value) if col("Precondition Value 2") else ""
        pre_field3 = _clean(ws.cell(row=r, column=col("Precondition Field 3")).value) if col("Precondition Field 3") else ""
        pre_op3 = _clean(ws.cell(row=r, column=col("Precondition Operator 3")).value) if col("Precondition Operator 3") else ""
        pre_val3 = _clean(ws.cell(row=r, column=col("Precondition Value 3")).value) if col("Precondition Value 3") else ""
        pre_logic = _clean(ws.cell(row=r, column=col("Precondition Logic")).value) if col("Precondition Logic") else ""
        expected = _clean(ws.cell(row=r, column=col("Expected Result")).value) if col("Expected Result") else ""
        cmp_table = _table_code(ws.cell(row=r, column=col("Compare Table")).value) if col("Compare Table") else ""
        cmp_field = _clean(ws.cell(row=r, column=col("Compare Field")).value) if col("Compare Field") else ""
        join_child = _normalize_pipe_fields(_clean(ws.cell(row=r, column=col("Join Child Fields")).value)) if col("Join Child Fields") else ""
        join_parent = _normalize_pipe_fields(_clean(ws.cell(row=r, column=col("Join Parent Fields")).value)) if col("Join Parent Fields") else ""
        # Aggregate columns (present only after the aggregate-type additions).
        groupby = _normalize_pipe_fields(_clean(ws.cell(row=r, column=col("GroupBy Fields")).value)) if col("GroupBy Fields") else ""
        cnt_filter_field = _clean(ws.cell(row=r, column=col("Count Filter Field")).value) if col("Count Filter Field") else ""
        cnt_filter_value = _normalize_pipe_fields(_clean(ws.cell(row=r, column=col("Count Filter Value")).value)) if col("Count Filter Value") else ""
        cnt_operator = _clean(ws.cell(row=r, column=col("Count Operator")).value) if col("Count Operator") else ""
        cnt_value = _clean(ws.cell(row=r, column=col("Count Value")).value) if col("Count Value") else ""
        # Two-hop (match_through_intermediate) columns.
        inter_table = _table_code(ws.cell(row=r, column=col("Intermediate Table")).value) if col("Intermediate Table") else ""
        inter_filter_field = _clean(ws.cell(row=r, column=col("Intermediate Filter Field")).value) if col("Intermediate Filter Field") else ""
        inter_filter_value = _clean(ws.cell(row=r, column=col("Intermediate Filter Value")).value) if col("Intermediate Filter Value") else ""
        join_inter = _normalize_pipe_fields(_clean(ws.cell(row=r, column=col("Join Intermediate Fields")).value)) if col("Join Intermediate Fields") else ""
        join_cmp = _normalize_pipe_fields(_clean(ws.cell(row=r, column=col("Join Compare Fields")).value)) if col("Join Compare Fields") else ""
        error_msg = _clean(ws.cell(row=r, column=col("Error Message")).value) if col("Error Message") else ""
        active_raw = _clean(ws.cell(row=r, column=col("Active")).value) if col("Active") else "Y"
        notes = _clean(ws.cell(row=r, column=col("Notes")).value) if col("Notes") else ""

        # Normalize 'active' to Y/N.
        active = "N" if active_raw.upper() in ("N", "NO", "FALSE", "0", "INACTIVE") else "Y"

        if rule_type not in evaluable:
            not_evaluated.append((xref_id, rule_type))

        rows.append({
            "xref_id": xref_id,
            "table_name": table_name,
            "target_field": target_field,
            "rule_type": rule_type,
            "severity": severity or "SEVERE",
            "precondition_field": pre_field,
            "precondition_operator": pre_op,
            "precondition_value": pre_val,
            "precondition_field_2": pre_field2,
            "precondition_operator_2": pre_op2,
            "precondition_value_2": pre_val2,
            "precondition_field_3": pre_field3,
            "precondition_operator_3": pre_op3,
            "precondition_value_3": pre_val3,
            "precondition_logic": pre_logic,
            "expected_result": expected.replace("\n", " ").replace("\r", " "),
            "compare_table": cmp_table,
            "compare_field": cmp_field,
            "join_child_fields": join_child,
            "join_parent_fields": join_parent,
            "groupby_fields": groupby,
            "count_filter_field": cnt_filter_field,
            "count_filter_value": cnt_filter_value,
            "count_operator": cnt_operator,
            "count_value": cnt_value,
            "intermediate_table": inter_table,
            "intermediate_filter_field": inter_filter_field,
            "intermediate_filter_value": inter_filter_value,
            "join_intermediate_fields": join_inter,
            "join_compare_fields": join_cmp,
            "error_message": error_msg.replace("\n", " ").replace("\r", " ")[:300],
            "active": active,
            "notes": notes.replace("\n", " ").replace("\r", " ")[:300],
        })

    return rows, stats, not_evaluated


# ---------------------------------------------------------------------------
# Writers + reconciliation
# ---------------------------------------------------------------------------

def _write_csv(path, columns, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def write_reconciliation_report(path, val_stats, val_unknown, xref_stats, xref_not_eval):
    lines = []
    lines.append("=" * 70)
    lines.append("VALIDATION_RULES_MASTER -> TEMPLATE RECONCILIATION REPORT")
    lines.append("=" * 70)

    lines.append("\n[FIELD-LEVEL VALIDATION RULES]  (validation_rules.csv)")
    total_rules = sum(s["rows"] for s in val_stats.values())
    total_fields = sum(s["fields"] for s in val_stats.values())
    for code, s in val_stats.items():
        lines.append(f"  {code}: {s['rows']} rules across {s['fields']} fields")
    lines.append(f"  TOTAL: {total_rules} rules across {total_fields} fields")

    if val_unknown:
        lines.append(f"\n  WARNING: {len(val_unknown)} rule(s) use a rule_type NOT in the engine registry")
        lines.append("  (written to CSV but will be SKIPPED at runtime until a handler exists):")
        for code, field, rt in val_unknown:
            lines.append(f"    - {code}.{field}: '{rt}'")
    else:
        lines.append("\n  OK: every field-level rule_type is registered in the engine.")

    lines.append("\n[CROSS-REFERENCE RULES]  (cross_references.csv)")
    total_xref = sum(xref_stats.values())
    for rt in sorted(xref_stats):
        registered = "registered" if rt in REGISTERED_CROSS_REFERENCE_TYPES else "NOT REGISTERED"
        lines.append(f"  {rt}: {xref_stats[rt]}  ({registered})")
    lines.append(f"  TOTAL: {total_xref} cross-reference rows")

    if xref_not_eval:
        lines.append(f"\n  NOTE: {len(xref_not_eval)} cross-reference(s) are carried forward but")
        lines.append("  the engine has no handler to EVALUATE them yet (preserved, not lost).")
        lines.append("  Add a handler + registry entry to activate — no adapter change needed:")
        # Summarize by type to keep the report short.
        by_type = {}
        for xid, rt in xref_not_eval:
            by_type.setdefault(rt, []).append(xid)
        for rt in sorted(by_type):
            ids = by_type[rt]
            sample = ", ".join(ids[:5]) + (" ..." if len(ids) > 5 else "")
            lines.append(f"    - {rt}: {len(ids)}  [{sample}]")
    else:
        lines.append("\n  OK: every cross-reference type has an engine handler.")

    report = "\n".join(lines) + "\n"
    with open(path, "w", encoding="utf-8") as f:
        f.write(report)
    print(report)


def main():
    if not os.path.exists(MASTER_PATH):
        raise FileNotFoundError(f"Master workbook not found: {MASTER_PATH}")
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    print(f"Reading {MASTER_PATH} ...")
    wb = openpyxl.load_workbook(MASTER_PATH, data_only=True)

    val_rows, val_stats, val_unknown = build_validation_rules(wb)
    xref_rows, xref_stats, xref_not_eval = build_cross_references(wb)
    wb.close()

    _write_csv(os.path.join(OUTPUT_DIR, "validation_rules.csv"),
               VALIDATION_RULES_COLUMNS, val_rows)
    _write_csv(os.path.join(OUTPUT_DIR, "cross_references.csv"),
               CROSS_REFERENCES_COLUMNS, xref_rows)
    _write_csv(os.path.join(OUTPUT_DIR, "table_dependencies.csv"),
               ["child_table", "parent_table", "join_fields_child", "join_fields_parent"],
               TABLE_DEPENDENCIES)

    write_reconciliation_report(
        os.path.join(OUTPUT_DIR, "_reconciliation_report.txt"),
        val_stats, val_unknown, xref_stats, xref_not_eval,
    )

    print("Output files written to:", OUTPUT_DIR)
    print("  - validation_rules.csv")
    print("  - cross_references.csv")
    print("  - table_dependencies.csv")
    print("  - _reconciliation_report.txt")


if __name__ == "__main__":
    main()
